"""JV-Data仕様書 (xlsx) の「フォーマット」シートからレコードレイアウトを生成する。

SDK 同梱の ``JV-Data仕様書_4.9.0.1.xlsx`` をそのまま読み、表 1 つを 1 レコード種別
として :class:`~jvstore.layout.RecordLayout` に変換する。手打ちのレイアウト定義を持たず
公式仕様書を唯一の出典にすることで、仕様改訂時は xlsx を差し替えて再生成するだけで済む。

シートの列構成（0 始まり）::

    1: 項番 / 表題    2: 枝番(a,b,c…)  3: キー   4: 項目名   5: 位置
    6: 繰返           7: バイト        8: 合計   9: 初期値  10: 説明
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Iterable, Iterator, Sequence

from .layout import DataSpec, Item, LayoutSet, RecordLayout

__all__ = ["parse_spec_workbook", "SpecParseError", "validate"]

# 「フォーマット」シートの列番号（0 始まり）。
(
    _COL_NO,
    _COL_SUB,
    _COL_KEY,
    _COL_NAME,
    _COL_POSITION,
    _COL_REPEAT,
    _COL_BYTES,
    _COL_TOTAL,
    _COL_DEFAULT,
    _COL_DESCRIPTION,
) = range(1, 11)

# 「データ種別一覧」シートの列番号（0 始まり）。
_COL_SPEC_NAME = 1
_COL_SPEC_ID = 2
_COL_SPEC_RECORD_ID = 5

_TITLE_RE = re.compile(r"^([０-９0-9]+)[．.]\s*(.+?)\s*$")
_RELATIVE_POSITION_RE = re.compile(r"^\(\s*(\d+)\s*\)$")
_RECORD_ID_RE = re.compile(r'["“]([A-Z0-9]{2})["”]')
_CATEGORY_RE = re.compile(
    r"（[１-３1-3]）\s*(蓄積系データ|速報系データ|セットアップデータ)"
)
_ZEN2HAN = str.maketrans("０１２３４５６７８９", "0123456789")
_INDENT_CHARS = "　 \t"

#: 繰返しブロックの見出し行は、項目名が ``<登録馬毎情報>`` のように囲まれている。
_GROUP_PREFIX = "<"

#: 「キー」列に付く印。仕様書の中で3種類の丸が混在している。
_KEY_MARKS = ("○", "◯", "〇")


class SpecParseError(RuntimeError):
    pass


def _text(cell: Any) -> str:
    return "" if cell is None else str(cell).strip()


def _as_int(cell: Any) -> int | None:
    """セルを整数にする。全角数字も読む。数値でなければ None。"""
    if cell is None:
        return None
    if isinstance(cell, int):
        return cell
    digits = str(cell).strip().translate(_ZEN2HAN)
    return int(digits) if digits.isdigit() else None


def parse_spec_workbook(xlsx: Path, sheet: str = "フォーマット") -> LayoutSet:
    """仕様書 xlsx を読み込んで :class:`LayoutSet` を返す。"""
    import openpyxl  # 生成時のみ必要（実行時依存にはしない）

    xlsx = Path(xlsx)
    workbook = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if sheet not in workbook.sheetnames:
        raise SpecParseError(f"シート {sheet!r} が見つかりません: {workbook.sheetnames}")
    rows = [list(row) for row in workbook[sheet].iter_rows(values_only=True)]
    layouts = [_parse_table(rows, begin, end) for begin, end in _table_ranges(rows)]
    version = re.search(r"(\d+(?:\.\d+)+)", xlsx.name)
    return LayoutSet(
        version=version.group(1) if version else "",
        source=xlsx.name,
        layouts={layout.record_id: layout for layout in layouts},
        dataspecs=_parse_dataspecs(workbook),
    )


def _parse_dataspecs(workbook: Any, sheet: str = "データ種別一覧") -> dict[str, DataSpec]:
    """「データ種別一覧」シートから データ種別ID → 含まれるレコード種別 を取り出す。

    列構成: 1=データ種別名 2=データ種別ID 3=フォーマットNo 4=レコード種別名
            5=レコード種別ID 6=収録内容
    データ種別ID セルは "DIFF DIFN" のように複数並ぶことがある。
    """
    if sheet not in workbook.sheetnames:
        return {}
    specs: dict[str, DataSpec] = {}
    category = ""
    current_ids: list[str] = []
    for row in workbook[sheet].iter_rows(values_only=True):
        cells = [_text(cell) for cell in row] + [""] * 8
        heading = _CATEGORY_RE.search(cells[_COL_SPEC_NAME])
        if heading:
            category = heading.group(1).replace("データ", "")
            current_ids = []
            continue
        if cells[_COL_SPEC_NAME] == "名称" or cells[_COL_SPEC_ID].startswith("データ"):
            continue  # 各セクションの見出し行
        spec_ids = cells[_COL_SPEC_ID].split()
        if spec_ids:
            current_ids = spec_ids
            _merge_specs(specs, spec_ids, cells[_COL_SPEC_NAME], category)
        record_id = cells[_COL_SPEC_RECORD_ID]
        if current_ids and re.fullmatch(r"[A-Z0-9]{2}", record_id):
            for spec_id in current_ids:
                record_ids = specs[spec_id].record_ids
                if record_id not in record_ids:
                    record_ids.append(record_id)
    return specs


def _merge_specs(
    specs: dict[str, DataSpec], spec_ids: Iterable[str], name: str, category: str
) -> None:
    """同じデータ種別IDが複数の区分に現れるので、名前と区分を足しこむ。"""
    for spec_id in spec_ids:
        spec = specs.setdefault(spec_id, DataSpec(spec_id, name, [], []))
        if not spec.name:
            spec.name = name
        if category and category not in spec.categories:
            spec.categories.append(category)


def _table_ranges(rows: list[list[Any]]) -> Iterator[tuple[int, int]]:
    """表題行のインデックスから、各表の [開始行, 終了行) を切り出す。"""
    titles = [index for index, row in enumerate(rows) if _is_title_row(row)]
    for position, begin in enumerate(titles):
        end = titles[position + 1] if position + 1 < len(titles) else len(rows)
        yield begin, end


def _is_title_row(row: Sequence[Any]) -> bool:
    if len(row) <= _COL_BYTES:
        return False
    if not _TITLE_RE.match(_text(row[_COL_NO])):
        return False
    # 「レコード長 / n / バイト」が同じ行に並ぶものだけを表題とみなす
    joined = " ".join(_text(cell) for cell in row)
    return "レコード長" in joined and _as_int(row[_COL_BYTES]) is not None


def _parse_table(rows: list[list[Any]], start: int, end: int) -> RecordLayout:
    index, title, length = _parse_title(rows[start])

    items: list[Item] = []
    # 繰返しブロックのスタック（仕様書 4.9 時点で入れ子は無いが将来に備える）
    open_groups: list[Item] = []
    record_id = ""

    for row in rows[start + 1 : end]:
        row = _padded_row(row)
        name = _item_name(row)
        size = _as_int(row[_COL_BYTES])
        if not _is_item_row(name, size):
            continue

        is_child = _as_int(row[_COL_NO]) is None  # 項番が数字でない行＝繰返しの内訳
        if not is_child:
            open_groups.clear()
        parent = open_groups[-1] if (is_child and open_groups) else None
        container = parent.children if parent is not None else items

        item = _build_item(row, name, size, parent, container)
        container.append(item)
        if name.startswith(_GROUP_PREFIX):
            open_groups.append(item)
        if not record_id:
            record_id = _record_id_of(item)

    if not record_id:
        raise SpecParseError(f"表 {title}: レコード種別IDを特定できません")
    return RecordLayout(
        index=index, title=title, record_id=record_id, length=length, items=items
    )


def _parse_title(row: Sequence[Any]) -> tuple[str, str, int]:
    """表題行から 表番号・表題・レコード長 を取り出す。"""
    matched = _TITLE_RE.match(_text(row[_COL_NO]))
    assert matched is not None  # _is_title_row を通った行しか渡らない
    title = matched.group(2)
    length = _as_int(row[_COL_BYTES])
    if length is None:
        raise SpecParseError(f"表 {title}: レコード長を読み取れません")
    return matched.group(1).translate(_ZEN2HAN), title, length


def _padded_row(row: Sequence[Any]) -> list[Any]:
    """説明列まで必ず存在するようにそろえる。末尾の空欄が省略されている行がある。"""
    padded = list(row)
    missing = _COL_DESCRIPTION + 1 - len(padded)
    return padded + [None] * missing if missing > 0 else padded


def _item_name(row: Sequence[Any]) -> str:
    raw = "" if row[_COL_NAME] is None else str(row[_COL_NAME])
    return raw.strip(_INDENT_CHARS).strip()


def _is_item_row(name: str, size: int | None) -> bool:
    """レイアウト上の実体を持つ行か。

    バイト数を持たない行は「<馬場別着回数>」のような見出し行や
    「連番002情報」のような説明専用行で、切り出す位置を持たない。
    """
    if not name or name == "項目名":
        return False
    return size is not None


def _build_item(
    row: Sequence[Any], name: str, size: int, parent: Item | None, container: list[Item]
) -> Item:
    no = _text(row[_COL_NO])
    sub = _text(row[_COL_SUB])
    return Item(
        no=(no + sub) if no else ((parent.no if parent else "") + sub),
        name=name.lstrip("<").rstrip(">"),
        offset=_item_position(row, container) - 1,
        size=size,
        repeat=_as_int(row[_COL_REPEAT]) or 1,
        is_key=_text(row[_COL_KEY]) in _KEY_MARKS,
        default=_text(row[_COL_DEFAULT]),
        comment=_text(row[_COL_DESCRIPTION]),
    )


def _item_position(row: Sequence[Any], container: list[Item]) -> int:
    """項目の開始位置（仕様書と同じ 1 始まり）。

    ブロック内の項目は "(3)" のように括弧つきの相対位置で書かれている。
    位置欄が空の行（重勝式の <重勝式対象レース情報> 内など）は直前の項目から連結する。
    """
    cell = row[_COL_POSITION]
    relative = _RELATIVE_POSITION_RE.match(_text(cell))
    if relative:
        return int(relative.group(1))
    absolute = _as_int(cell)
    if absolute is not None:
        return absolute
    previous = container[-1] if container else None
    return (previous.offset + previous.total + 1) if previous else 1


def _record_id_of(item: Item) -> str:
    """「レコード種別ID」項目の説明欄に書かれた ``"RA"`` を取り出す。"""
    if item.name != "レコード種別ID":
        return ""
    found = _RECORD_ID_RE.search(item.comment)
    return found.group(1) if found else ""


def validate(layout: RecordLayout) -> list[str]:
    """レイアウトの自己整合性を検査し、問題点の一覧を返す（空なら健全）。"""
    problems: list[str] = []
    _check_container(
        layout.items, f"{layout.record_id}({layout.title})", layout.length, problems
    )
    return problems


def _check_container(
    items: Iterable[Item], container: str, container_size: int, problems: list[str]
) -> None:
    """項目が隙間なく並び、合計がコンテナの大きさと一致することを確かめる。"""
    cursor = 0
    for item in items:
        if item.offset != cursor:
            problems.append(
                f"{container}: {item.no} {item.name} の位置が不連続 "
                f"(期待 {cursor + 1}, 実際 {item.offset + 1})"
            )
        cursor = item.offset + item.total
        if item.children:
            inner = sum(child.total for child in item.children)
            if inner != item.size:
                problems.append(
                    f"{container}: <{item.name}> の内訳合計 {inner} が "
                    f"1繰返し分 {item.size} と不一致"
                )
            _check_container(
                item.children, f"{container}/<{item.name}>", item.size, problems
            )
    if cursor != container_size:
        problems.append(
            f"{container}: 合計 {cursor} バイトがレコード長/ブロック長 "
            f"{container_size} と不一致"
        )
