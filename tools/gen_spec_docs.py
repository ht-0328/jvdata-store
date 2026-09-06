r"""JV-Data仕様書から、参照用のドキュメントを生成する。

出力先は ``docs/reference/``。手で書き足さないこと。仕様が改訂されたら
``jvstore spec build`` で ``layouts.json`` を作り直し、このスクリプトを流し直す。

    uv run python tools/gen_spec_docs.py --xlsx "...\JV-Data仕様書_4.9.0.1.xlsx"

``--xlsx`` を省くとコード表は作らず、レイアウトだけを生成する
（レイアウトは同梱の ``layouts.json`` から作るので xlsx が要らない）。
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from jvstore.layout import Item, RecordLayout, load_layouts  # noqa: E402

OUT = ROOT / "docs" / "reference"

#: 速報系専用の表。蓄積系（過去N年の取得）には含まれない。
REALTIME_ONLY = {"WH", "WE", "AV", "JC", "TC", "CC"}


def esc(text: str) -> str:
    """表のセルに入れられる形にする。改行と縦棒を潰す。"""
    return (text or "").replace("|", "\\|").replace("\n", " ").strip()


def cell(text: str) -> str:
    """本文セル用。仕様書の説明には ``[抽](指定)`` のように、そのまま置くと
    Markdown のリンクと解釈される並びがある。角括弧を打ち消しておく。"""
    return esc(text).replace("[", "&#91;").replace("]", "&#93;")


def code_links(comment: str) -> str:
    """`<コード表 2001.競馬場コード>参照` をコード表への相対リンクにする。"""
    return re.sub(
        r"<コード表\s*(\d{4})\.([^>]+)>",
        lambda m: f"[コード表 {m.group(1)}.{m.group(2)}](codes.md#{m.group(1)})",
        comment,
    )


# ---------------------------------------------------------------- レコード


def item_rows(items: list[Item], depth: int = 0) -> list[str]:
    """項目を1行ずつ Markdown の表の行にする。繰返しブロックは字下げして続ける。"""
    out: list[str] = []
    for it in items:
        indent = "　" * depth + ("└ " if depth else "")
        pos = it.offset + 1 if depth == 0 else f"(+{it.offset})"
        rep = f"×{it.repeat}" if it.repeat > 1 else ""
        out.append(
            f"| {it.no} | {'○' if it.is_key else ''} | {indent}{cell(it.name)} | "
            f"{pos} | {it.size} | {rep} | {esc(it.default)} | {code_links(cell(it.comment))} |"
        )
        if it.children:
            out.extend(item_rows(it.children, depth + 1))
    return out


def record_page(lay: RecordLayout) -> str:
    keys = [i.name for i in lay.items if i.is_key and not i.children]
    blocks = [i for i in lay.items if i.children]
    plain_reps = [i for i in lay.items if i.repeat > 1 and not i.children]
    kind = "速報系のみ" if lay.record_id in REALTIME_ONLY else "蓄積系"

    lines = [
        f"# {lay.record_id} — {lay.title}",
        "",
        f"| | |",
        f"| :--- | :--- |",
        f"| レコード種別ID | `{lay.record_id}` |",
        f"| 表番号 | {lay.index} |",
        f"| レコード長 | {lay.length:,} バイト |",
        f"| 項目数 | {len(lay.items)} |",
        f"| 区分 | {kind} |",
        f"| DuckDB のテーブル名 | `{lay.record_id.lower()}` |",
        "",
        "## キー",
        "",
    ]
    if keys:
        lines += [
            "この組み合わせで1件が決まります（仕様書の「キー」列）。",
            "",
            *[f"{n}. {k}" for n, k in enumerate(keys, 1)],
        ]
    else:
        lines.append("仕様書にキーの指定がありません。")
    lines.append("")

    if blocks:
        lines += [
            "## 繰返しブロック",
            "",
            "本ツールでは、繰返しブロックは**子テーブルに分けます**"
            "（横に並べると列数が扱えなくなるため）。",
            "",
            "| ブロック | 繰返し | 1回分のバイト数 | 子テーブル名 |",
            "| :--- | ---: | ---: | :--- |",
        ]
        for b in blocks:
            slug = re.sub(r"[^0-9A-Za-z぀-ヿ一-鿿]+", "_", b.name).strip("_")
            lines.append(
                f"| {cell(b.name)} | {b.repeat} | {b.size // b.repeat if b.repeat else b.size} "
                f"| `{lay.record_id.lower()}__{slug}` |"
            )
        lines.append("")

    if plain_reps:
        lines += [
            "## 連番になる項目",
            "",
            "内訳を持たない繰返し項目は、子テーブルにはせず**連番付きの列**に展開します。",
            "`本賞金` が7回繰返しなら `本賞金_1` 〜 `本賞金_7` になります。",
            "",
            "| 項目 | 繰返し | 展開後の列名 |",
            "| :--- | ---: | :--- |",
        ]
        for r in plain_reps:
            w = len(str(r.repeat))
            first = f"{esc(r.name)}_{1:0{w}d}"
            last = f"{esc(r.name)}_{r.repeat:0{w}d}"
            lines.append(f"| {cell(r.name)} | {r.repeat} | `{first}` 〜 `{last}` |")
        lines.append("")

    lines += [
        "## 全項目",
        "",
        "「位置」は1始まりのバイト位置です。繰返しブロックの中の項目は、"
        "ブロック先頭からの相対位置を `(+n)` で示します（0始まり）。",
        "",
        "| 項番 | キー | 項目名 | 位置 | バイト | 繰返 | 初期値 | 説明 |",
        "| ---: | :---: | :--- | ---: | ---: | ---: | :--- | :--- |",
        *item_rows(lay.items),
        "",
    ]
    return "\n".join(lines)


def index_page(layouts) -> str:
    lines = [
        "# レコード索引",
        "",
        "JV-Data の全38レコードです。`JV-Data仕様書_4.9.0.1.xlsx` から生成しています。",
        "",
        "## 蓄積系（過去N年ぶんの取得対象）",
        "",
        "| 表番号 | ID | 名前 | レコード長 | 項目数 | キー |",
        "| ---: | :--- | :--- | ---: | ---: | :--- |",
    ]
    rt: list[str] = []
    for lay in layouts:
        keys = "、".join(i.name for i in lay.items if i.is_key and not i.children) or "—"
        row = (
            f"| {lay.index} | [{lay.record_id}]({lay.record_id}.md) | {lay.title} "
            f"| {lay.length:,} | {len(lay.items)} | {keys} |"
        )
        (rt if lay.record_id in REALTIME_ONLY else lines).append(row)
    lines += [
        "",
        "## 速報系のみ",
        "",
        "サーバでの提供期間が1週間しかないため、過去N年ぶんの取得には含まれません。",
        "`jvstore rt` で当日ぶんを取ります。",
        "",
        "| 表番号 | ID | 名前 | レコード長 | 項目数 | キー |",
        "| ---: | :--- | :--- | ---: | ---: | :--- |",
        *rt,
        "",
    ]
    return "\n".join(lines)


# ---------------------------------------------------------------- コード表


def codes_page(xlsx: Path) -> str:
    from openpyxl import load_workbook

    ws = load_workbook(xlsx, read_only=True, data_only=True)["コード表"]
    rows = [
        [("" if c is None else str(c).strip()) for c in r]
        for r in ws.iter_rows(values_only=True)
    ]

    starts = [i for i, r in enumerate(rows) if len(r) > 1 and re.match(r"^\d{4}\.", r[1])]
    out = [
        "# コード表",
        "",
        "`JV-Data仕様書_4.9.0.1.xlsx` の「コード表」シートから生成しています。",
        "",
        "各項目の説明にある `<コード表 2001.競馬場コード>参照` は、この表を指しています。",
        "",
        "> [!NOTE]",
        "> **コードは文字列として持ちます**",
        ">",
        "> 本ツールは値を仕様書の桁のまま保存します。`競馬場コード` は `01` であって",
        "> `1` ではありません。先頭のゼロを落とすと、この表と突き合わせられなくなります。",
        "",
        "## 目次",
        "",
    ]
    titles = [rows[i][1] for i in starts]
    for t in titles:
        num = t.split(".")[0]
        out.append(f"- [{t}](#{num})")
    out.append("")

    for n, start in enumerate(starts):
        end = starts[n + 1] if n + 1 < len(starts) else len(rows)
        title = rows[start][1]
        num = title.split(".")[0]
        body = rows[start + 1 : end]
        # 1行目が「バイト数/値/内容」の見出し、2行目が内容の内訳。
        sub = body[1] if len(body) > 1 else []
        heads = [h for h in sub[3:] if h] or ["内容"]
        used = [j for j, h in enumerate(sub[3:], start=3) if h] or [3]
        size = next((r[1] for r in body[2:] if len(r) > 1 and r[1]), "")

        out += [
            f'## {title} {{: #{num} }}',
            "",
            f"バイト数: {size}" if size else "",
            "",
            "| 値 | " + " | ".join(cell(h) for h in heads) + " |",
            "| :--- | " + " | ".join(":---" for _ in heads) + " |",
        ]
        for r in body[2:]:
            if len(r) < 3 or not r[2]:
                continue
            cells = [cell(r[j]) if j < len(r) else "" for j in used]
            out.append(f"| `{esc(r[2])}` | " + " | ".join(cells) + " |")
        out.append("")
    return "\n".join(x for x in out if x is not None)


# ---------------------------------------------------------------- 実行


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--xlsx", help="JV-Data仕様書 xlsx。省くとコード表を作らない")
    args = ap.parse_args()

    OUT.mkdir(parents=True, exist_ok=True)
    layouts = load_layouts()

    (OUT / "index.md").write_text(index_page(layouts), encoding="utf-8")
    n = 0
    for lay in layouts:
        (OUT / f"{lay.record_id}.md").write_text(record_page(lay), encoding="utf-8")
        n += 1
    print(f"レコード {n} 件 -> {OUT}")

    if args.xlsx:
        (OUT / "codes.md").write_text(codes_page(Path(args.xlsx)), encoding="utf-8")
        print(f"コード表 -> {OUT / 'codes.md'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
